#!/usr/bin/env python3
"""
从 PDF 和 DOCX 文件中提取文本。
用法: python3 extract_text.py <file_path> [file_path ...]
输出: 各文件的文本内容，以 ===FILE=== 分隔
"""
import sys
import os

def extract_pdf(path: str) -> str:
    """用 pdfplumber 提取 PDF 文本；图片型 PDF 降级到 OCR。"""
    try:
        import pdfplumber
        texts = []
        with pdfplumber.open(path) as pdf:
            for page in pdf.pages:
                t = page.extract_text()
                if t:
                    texts.append(t)
        content = "\n".join(texts).strip()
        # 如果文字提取结果极少（<50字符），说明是扫描件
        if len(content) < 50:
            return ocr_pdf(path)
        return content
    except Exception as e:
        return f"[PDF读取失败: {e}]\n{ocr_pdf(path)}"


def ocr_pdf(path: str) -> str:
    """对图片型 PDF 使用 pytesseract OCR。"""
    try:
        from pdf2image import convert_from_path
        import pytesseract
        images = convert_from_path(path, dpi=200)
        texts = []
        for img in images:
            t = pytesseract.image_to_string(img, lang='chi_sim+eng')
            texts.append(t)
        return "\n".join(texts)
    except Exception as e:
        return f"[OCR失败: {e}]"


def extract_docx(path: str) -> str:
    """用 python-docx 提取 DOCX 文本。"""
    try:
        from docx import Document
        doc = Document(path)
        paragraphs = []
        for para in doc.paragraphs:
            if para.text.strip():
                paragraphs.append(para.text.strip())
        # 提取表格
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    paragraphs.append(" | ".join(cells))
        return "\n".join(paragraphs)
    except Exception as e:
        return f"[DOCX读取失败: {e}]"


def extract_file(path: str) -> str:
    ext = os.path.splitext(path)[1].lower()
    if ext == ".pdf":
        return extract_pdf(path)
    elif ext in (".docx", ".doc"):
        return extract_docx(path)
    else:
        return f"[不支持的格式: {ext}]"


def main():
    if len(sys.argv) < 2:
        print("用法: python3 extract_text.py <file_path> [file_path ...]", file=sys.stderr)
        sys.exit(1)

    results = []
    for fp in sys.argv[1:]:
        if not os.path.exists(fp):
            results.append(f"=== {os.path.basename(fp)} ===\n[文件不存在]\n")
            continue
        text = extract_file(fp)
        results.append(f"=== {os.path.basename(fp)} ===\n{text}\n")

    print("\n".join(results))


if __name__ == "__main__":
    main()
